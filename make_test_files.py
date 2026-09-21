"""
Generates sample CV files in several layouts and records what the ORIGINAL
backend (pdfplumber + python-docx + ai_engine) extracts from each, so that
tests/files.test.mjs can compare the browser extraction against it.

Layouts covered:
  single   - ordinary single-column CV
  twocol   - two text columns, no drawn shapes (pdfplumber column-split path: left column, then right column)
  table    - experience laid out as a ruled table (pdfplumber "has tables" path)
  sidebar  - two columns with a coloured sidebar rectangle (rect => treated as table)
  docx     - Word document with paragraphs and a table (tables are ignored by python-docx)
  txt      - plain text

Usage: python3 tests/make_test_files.py <old-backend-dir> <out-dir>
"""
import json
import os
import sys
from dataclasses import asdict

backend_dir, out_dir = sys.argv[1], sys.argv[2]
sys.path.insert(0, backend_dir)
os.makedirs(out_dir, exist_ok=True)
import ai_engine  # noqa: E402

from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.pdfgen import canvas  # noqa: E402
from reportlab.lib import colors  # noqa: E402
import docx  # noqa: E402

PEOPLE = [
    dict(name="Budi Santoso", email="budi.santoso@example.com", phone="081234567890",
         edu="S1 Teknik Industri, Universitas Indonesia, 2010 - 2014",
         jobs=[("Production Supervisor", "PT Contoh Sejahtera", "2016 - 2020"),
               ("Production Manager", "PT Global Industries", "2020 - Present")],
         skills="Six Sigma, Lean Manufacturing, Production Planning, Quality Control, Leadership, Communication",
         certs=["Six Sigma Green Belt 2018", "ISO 9001 Lead Auditor 2019"]),
    dict(name="Siti Rahayu", email="siti.rahayu@mail.co.id", phone="08129876543",
         edu="D3 Akuntansi, Politeknik Negeri Jakarta, 2015 - 2018",
         jobs=[("Finance Staff", "PT Maju Bersama", "2018 - 2022"),
               ("Senior Accountant", "CV Sentosa Abadi", "2022 - sekarang")],
         skills="Accounting, Excel, Taxation, Budgeting, Financial Analysis, komunikasi",
         certs=["Brevet A dan B 2019"]),
    dict(name="Rudi Hartono", email="rudi.hartono@site.org", phone="085678901234",
         edu="S2 Manajemen, Universitas Gadjah Mada, 2012 - 2014",
         jobs=[("Head of Sales", "PT Alpha Tbk", "2014 - 2019"),
               ("Sales Director", "PT Beta Group", "2019 - Present"),
               ("Sales Executive", "CV Gamma", "2010 - 2014")],
         skills="Sales Forecasting, CRM, Salesforce, Negotiation, Leadership, Public Speaking",
         certs=["PMP 2016", "TOEFL 2015"]),
]


def lines_for(p, wrap=True):
    out = [p["name"], f"Email: {p['email']}  Telp: {p['phone']}", "", "PENDIDIKAN", p["edu"], "", "PENGALAMAN KERJA"]
    for pos, comp, yrs in p["jobs"]:
        out.append(f"{pos}, {comp}, {yrs}")
    out += ["", "KEAHLIAN", p["skills"], "", "SERTIFIKASI"] + p["certs"]
    return out


def make_single(path, p):
    c = canvas.Canvas(path, pagesize=A4)
    y = 800
    c.setFont("Helvetica-Bold", 16)
    c.drawString(60, y, p["name"]); y -= 24
    c.setFont("Helvetica", 10)
    for ln in lines_for(p)[1:]:
        c.drawString(60, y, ln); y -= 15
    c.save()


def make_twocol(path, p):
    c = canvas.Canvas(path, pagesize=A4)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(60, 800, p["name"])
    c.setFont("Helvetica", 10)
    # left column: contact/education/skills
    left = [f"Email: {p['email']}", f"Telp: {p['phone']}", "", "PENDIDIKAN", p["edu"][:34], p["edu"][34:].strip(), "", "KEAHLIAN"]
    skills = [s.strip() for s in p["skills"].split(",")]
    left += skills
    y = 770
    for ln in left:
        c.drawString(50, y, ln); y -= 14
    # right column: experience/certs
    right = ["PENGALAMAN KERJA"]
    for pos, comp, yrs in p["jobs"]:
        right += [f"{pos}, {comp}", yrs, ""]
    right += ["SERTIFIKASI"] + p["certs"]
    y = 770
    for ln in right:
        c.drawString(350, y, ln); y -= 14
    c.save()


def make_table(path, p):
    c = canvas.Canvas(path, pagesize=A4)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(60, 800, p["name"])
    c.setFont("Helvetica", 10)
    c.drawString(60, 780, f"Email: {p['email']}  Telp: {p['phone']}")
    c.drawString(60, 760, "PENDIDIKAN")
    c.drawString(60, 745, p["edu"])
    c.drawString(60, 720, "PENGALAMAN KERJA")
    # ruled table
    x0, x1, x2, x3 = 60, 200, 380, 500
    y = 705
    rows = [("Posisi", "Perusahaan", "Tahun")] + [(pos, comp, yrs) for pos, comp, yrs in p["jobs"]]
    top = y + 12
    for r in rows:
        c.drawString(x0 + 4, y, r[0]); c.drawString(x1 + 4, y, r[1]); c.drawString(x2 + 4, y, r[2])
        y -= 18
    bottom = y + 12
    c.setStrokeColor(colors.black)
    for yy in [top - 18 * i for i in range(len(rows) + 1)]:
        c.line(x0, yy, x3, yy)
    for xx in (x0, x1, x2, x3):
        c.line(xx, top, xx, bottom)
    c.drawString(60, y - 10, "KEAHLIAN")
    c.drawString(60, y - 25, p["skills"][:90])
    c.drawString(60, y - 50, "SERTIFIKASI")
    yy = y - 65
    for ce in p["certs"]:
        c.drawString(60, yy, ce); yy -= 14
    c.save()


def make_sidebar(path, p):
    c = canvas.Canvas(path, pagesize=A4)
    c.setFillColor(colors.HexColor("#DDE6F0"))
    c.rect(0, 0, 200, 842, stroke=0, fill=1)
    c.setFillColor(colors.black)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(225, 800, p["name"])
    c.setFont("Helvetica", 10)
    y = 770
    for ln in [f"Email: {p['email']}", f"Telp: {p['phone']}", "", "KEAHLIAN"] + [s.strip() for s in p["skills"].split(",")]:
        c.drawString(20, y, ln); y -= 14
    y = 770
    right = ["PENDIDIKAN", p["edu"][:40], p["edu"][40:].strip(), "", "PENGALAMAN KERJA"]
    for pos, comp, yrs in p["jobs"]:
        right += [f"{pos}, {comp}, {yrs}"]
    right += ["", "SERTIFIKASI"] + p["certs"]
    for ln in right:
        c.drawString(225, y, ln); y -= 14
    c.save()


def make_docx(path, p):
    d = docx.Document()
    d.add_paragraph(p["name"])
    d.add_paragraph(f"Email: {p['email']}  Telp: {p['phone']}")
    d.add_paragraph("PENDIDIKAN")
    d.add_paragraph(p["edu"])
    d.add_paragraph("PENGALAMAN KERJA")
    for pos, comp, yrs in p["jobs"]:
        d.add_paragraph(f"{pos}, {comp}, {yrs}")
    d.add_paragraph("KEAHLIAN")
    d.add_paragraph(p["skills"])
    d.add_paragraph("SERTIFIKASI")
    for ce in p["certs"]:
        para = d.add_paragraph()
        para.add_run(ce.split()[0] + " ")
        para.add_run(" ".join(ce.split()[1:]))
    tbl = d.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Ignored in table"
    tbl.cell(0, 1).text = "Kolom kedua 1999 - 2001"
    d.save(path)


def make_txt(path, p):
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines_for(p)))


records = []
makers = [("single", make_single, ".pdf"), ("twocol", make_twocol, ".pdf"), ("table", make_table, ".pdf"),
          ("sidebar", make_sidebar, ".pdf"), ("docx", make_docx, ".docx"), ("txt", make_txt, ".txt")]
for i, p in enumerate(PEOPLE):
    for kind, fn, ext in makers:
        fname = f"cv_{kind}_{i}{ext}"
        path = os.path.join(out_dir, fname)
        fn(path, p)
        text = ai_engine.extract_text_from_file(path)
        prof = asdict(ai_engine.extract_candidate_profile(text))
        records.append({"file": fname, "kind": kind, "text": text, "profile": prof,
                        "sha256": ai_engine.compute_file_hash(path)})

json.dump(records, open(os.path.join(out_dir, "expected.json"), "w"), ensure_ascii=False, indent=1)
print(f"wrote {len(records)} files to {out_dir}")
